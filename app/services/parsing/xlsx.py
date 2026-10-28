"""Excel 解析：openpyxl 只读模式逐 sheet 抽表格。

企业知识库的常见形态是制度表、花名册、报销标准表，这些都在 Excel 里。
每个 sheet 当作一个章节（section_path = "工作表: <sheet名>"），
整张表转成 Markdown 表格块，让大模型能读懂行列关系。

只读模式（read_only=True）流式读，不把整个文件加载进内存；
data_only=True 取公式计算后的值，而不是公式串本身。
"""

from openpyxl import load_workbook

from app.core.logging import logger
from app.services.parsing.base import Block, ParseResult, table_to_markdown


def _sheet_to_rows(ws) -> list[list[str]]:
    """把一个 sheet 读成二维字符串数组，跳过全空行。"""
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        # 整行都是空单元格就跳过
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        rows.append(["" if c is None else str(c).strip() for c in row])
    return rows


def parse(source) -> ParseResult:
    blocks: list[Block] = []
    # source 可能是路径或 BytesIO（UploadFile 包过来的），openpyxl 都认
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet_count = 0
        for ws in wb.worksheets:
            sheet_count += 1
            sheet_name = ws.title or f"Sheet{sheet_count}"
            rows = _sheet_to_rows(ws)
            if not rows:
                logger.info("Excel sheet 为空，跳过 sheet=%s", sheet_name)
                continue
            md = table_to_markdown(rows)
            if md:
                # page_no 对 Excel 无意义，置 None；章节用工作表名
                blocks.append(Block(md, None, f"工作表: {sheet_name}", "table"))
    finally:
        wb.close()

    return ParseResult(
        blocks=blocks,
        source_type="xlsx",
        page_count=sheet_count,
        needs_ocr=False,
    )
