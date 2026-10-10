"""Word 解析：python-docx 按正文顺序遍历段落与表格。

Word 的标题带样式标记（Heading 1/2/3），比 PDF 的启发式猜测准确得多，
所以这里的章节路径是真实层级，不是猜的。
"""

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.parsing.base import (
    Block,
    ParseResult,
    SectionTracker,
    table_to_markdown,
)


def _iter_body(document):
    """按正文顺序同时遍历段落和表格。

    python-docx 的 document.paragraphs 和 document.tables 是两个独立列表，
    分别取会丢失「表格在文档中的位置」。这里直接遍历 XML 子节点来保持顺序。
    """
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    """从段落样式名推断标题层级，如 'Heading 2' → 2。非标题返回 None。"""
    style = paragraph.style
    name = (style.name or "") if style is not None else ""
    if not name.lower().startswith("heading"):
        return None
    try:
        return int(name.split()[-1])
    except (ValueError, IndexError):
        return 1


def parse(source) -> ParseResult:
    document = Document(source)
    tracker = SectionTracker()
    blocks: list[Block] = []

    for item in _iter_body(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item)
            if level is not None:
                tracker.push(level, text)
                blocks.append(Block(text, None, tracker.path, "heading"))
            else:
                blocks.append(Block(text, None, tracker.path, "text"))
        elif isinstance(item, Table):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            md = table_to_markdown(rows)
            if md:
                blocks.append(Block(md, None, tracker.path, "table"))

    return ParseResult(blocks=blocks, source_type="docx", page_count=None)
